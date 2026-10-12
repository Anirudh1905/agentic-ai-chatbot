from pathlib import Path

import docx
import pymupdf

from chatbot.ingestion.chunking import chunk_document
from chatbot.ingestion.loaders import is_supported, load
from chatbot.ingestion.pipeline import discover

MARKDOWN = """# Refund Policy

Intro paragraph.

## Standard Returns

Return within 30 days.

## Sale Items

### Clearance

Clearance items are final sale.
"""


def test_markdown_chunks_carry_heading_path_without_title(tmp_path: Path) -> None:
    path = tmp_path / "refunds.md"
    path.write_text(MARKDOWN)

    doc = load(path)
    chunks = chunk_document(doc, chunk_size=200, chunk_overlap=0)

    assert doc.title == "Refund Policy"
    assert [c.heading for c in chunks] == [None, "Standard Returns", "Sale Items > Clearance"]
    assert chunks[2].content == "Clearance items are final sale."
    assert chunks[2].embed_text.startswith("Refund Policy\nSale Items > Clearance\n")
    assert [c.index for c in chunks] == [0, 1, 2]


def test_long_sections_split_with_overlap(tmp_path: Path) -> None:
    path = tmp_path / "long.md"
    path.write_text("# Long\n\n" + " ".join(f"word{i}" for i in range(400)))

    chunks = chunk_document(load(path), chunk_size=100, chunk_overlap=20)

    assert len(chunks) > 3
    assert len({c.content_hash for c in chunks}) == len(chunks)


def test_content_hash_ignores_whitespace_changes(tmp_path: Path) -> None:
    a, b = tmp_path / "a.md", tmp_path / "b.md"
    a.write_text("# Doc\n\nSame   text here.")
    b.write_text("# Doc\n\nSame text\nhere.")

    assert (
        chunk_document(load(a), 200, 0)[0].content_hash
        == chunk_document(load(b), 200, 0)[0].content_hash
    )


def test_html_loader_drops_head_and_keeps_headings(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>Ignored</title><style>.x{}</style></head>"
        "<body><h1>Warranty</h1><p>Lifetime cover.</p>"
        "<h2>Claims</h2><p>Send a photo.</p></body></html>"
    )

    doc = load(path)

    assert doc.title == "Warranty"
    assert "Ignored" not in doc.pages[0].text
    assert "## Claims" in doc.pages[0].text


def test_docx_loader_maps_heading_styles(tmp_path: Path) -> None:
    path = tmp_path / "handbook.docx"
    d = docx.Document()
    d.add_heading("Handbook", level=1)
    d.add_heading("PTO", level=2)
    d.add_paragraph("Twenty days per year.")
    d.save(str(path))

    text = load(path).pages[0].text

    assert text.splitlines()[0] == "# Handbook"
    assert "## PTO" in text
    assert "Twenty days per year." in text


def test_pdf_loader_keeps_page_numbers(tmp_path: Path) -> None:
    path = tmp_path / "doc.pdf"
    pdf = pymupdf.open()
    for text in ["First page text", "Second page text"]:
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(path)

    pages = load(path).pages

    assert [p.number for p in pages] == [1, 2]
    assert "Second page text" in pages[1].text


def test_discover_finds_supported_files_only(tmp_path: Path) -> None:
    (tmp_path / "sub").mkdir()
    for name in ["a.md", "sub/b.txt", "c.png", "d.HTML"]:
        (tmp_path / name).write_text("x")

    found = [p.relative_to(tmp_path).as_posix() for p in discover(tmp_path)]

    assert found == ["a.md", "d.HTML", "sub/b.txt"]
    assert not is_supported(tmp_path / "c.png")
